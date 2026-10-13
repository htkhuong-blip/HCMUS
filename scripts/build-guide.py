"""Generate editable Word and Vietnamese PDF from the maintained Markdown guide.
Optional document build only: pip install -r docs/requirements.txt.
"""
from pathlib import Path
from html import escape
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
ROOT = Path(__file__).resolve().parents[1]
source = ROOT / 'docs/thuyet-minh-thuc-hien.md'
lines = source.read_text().splitlines()
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Cm(2)
section.left_margin = section.right_margin = Cm(2.2)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
for heading in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    doc.styles[heading].font.color.rgb = RGBColor.from_string('245B4B')
section.footer.paragraphs[0].text = 'HCMUS HeritageKit 0.2 • Thuyết minh quy trình • Không thay thế hồ sơ xác nhận thực tế'
for line in lines:
    if not line.strip():
        continue
    if line.startswith('# '):
        doc.add_heading(line[2:], 0)
    elif line.startswith('## '):
        doc.add_heading(line[3:], 1)
    elif line.startswith('### '):
        doc.add_heading(line[4:], 2)
    else:
        doc.add_paragraph(line)
doc.save(ROOT / 'docs/thuyet-minh-thuc-hien.docx')
font = Path('/usr/share/fonts/truetype/dejavu')
pdfmetrics.registerFont(TTFont('VietSans', str(font / 'DejaVuSans.ttf')))
pdfmetrics.registerFont(TTFont('VietSansBold', str(font / 'DejaVuSans-Bold.ttf')))
styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name='VNBody', fontName='VietSans', fontSize=9.5, leading=15, spaceAfter=10))
styles.add(ParagraphStyle(name='VNTitle', fontName='VietSansBold', fontSize=22, leading=29, textColor=colors.HexColor('#245b4b'), spaceAfter=20))
styles.add(ParagraphStyle(name='VNHeading', fontName='VietSansBold', fontSize=14, leading=20, spaceBefore=14, spaceAfter=10, textColor=colors.HexColor('#245b4b'), keepWithNext=True))
styles.add(ParagraphStyle(name='VNSub', fontName='VietSansBold', fontSize=11, leading=17, spaceBefore=10, spaceAfter=7, keepWithNext=True))
story = []
for line in lines:
    if not line.strip():
        continue
    level = len(line) - len(line.lstrip('#'))
    name = 'VNTitle' if level == 1 else 'VNHeading' if level == 2 else 'VNSub' if level == 3 else 'VNBody'
    story.append(Paragraph(escape(line.lstrip('# ').strip()), styles[name]))
def footer(canvas, document):
    canvas.saveState()
    canvas.setFont('VietSans', 8)
    canvas.setFillColor(colors.HexColor('#637469'))
    canvas.drawString(45, 24, 'HCMUS HeritageKit • Thuyết minh cách thực hiện')
    canvas.drawRightString(A4[0]-45, 24, str(document.page))
    canvas.restoreState()
SimpleDocTemplate(str(ROOT / 'docs/thuyet-minh-thuc-hien.pdf'), pagesize=A4, leftMargin=45, rightMargin=45, topMargin=40, bottomMargin=45, title='HCMUS HeritageKit - Thuyết minh cách thực hiện', author='HCMUS HeritageKit').build(story, onFirstPage=footer, onLaterPages=footer)
print('Generated docs/thuyet-minh-thuc-hien.docx and .pdf')
